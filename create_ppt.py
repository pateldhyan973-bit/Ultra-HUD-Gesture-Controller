import os
from pptx import Presentation
from pptx.util import Inches

def create_presentation():
    prs = Presentation()

    # Slide 1: Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    title = slide1.shapes.title
    subtitle = slide1.placeholders[1]
    title.text = "Gesture Desktop Controller"
    subtitle.text = "Hands-free PC control using AI\nCreated with OpenCV and MediaPipe"

    # Slide 2: Features
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide2.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Key Features"

    tf = body_shape.text_frame
    tf.text = "Mouse Movement: Move cursor by moving your index finger"
    
    p = tf.add_paragraph()
    p.text = "Clicking: Pinch index and thumb to left-click or drag"
    
    p = tf.add_paragraph()
    p.text = "Right Click: Pinch middle finger and thumb"
    
    p = tf.add_paragraph()
    p.text = "Scrolling: Move two fingers up or down"

    p = tf.add_paragraph()
    p.text = "Volume Control: Pinch all fingers and adjust distance"

    # Slide 3: Video Demo
    blank_slide_layout = prs.slide_layouts[5] # Title only
    slide3 = prs.slides.add_slide(blank_slide_layout)
    slide3.shapes.title.text = "Live Demonstration"

    video_path = os.path.join('test_output', 'gesture_recording.mp4')
    
    if os.path.exists(video_path):
        # We need a poster image for the video (we can use one of the frames)
        poster_image = os.path.join('test_output', 'frame_4.png')
        if not os.path.exists(poster_image):
            poster_image = None
            
        left = Inches(1.5)
        top = Inches(2.0)
        width = Inches(7)
        height = Inches(5)
        
        try:
            slide3.shapes.add_movie(video_path, left, top, width, height, poster_frame_image=poster_image, mime_type='video/mp4')
        except Exception as e:
            print(f"Error adding video: {e}")
            tf = slide3.shapes.add_textbox(left, top, width, Inches(1)).text_frame
            tf.text = f"Video could not be embedded. File is at: {video_path}"
    else:
        left = Inches(1)
        top = Inches(3)
        width = Inches(8)
        height = Inches(1)
        tf = slide3.shapes.add_textbox(left, top, width, height).text_frame
        tf.text = "Video file not found at test_output/gesture_recording.mp4"

    # Save presentation
    output_ppt = "Gesture_Controller_Presentation.pptx"
    prs.save(output_ppt)
    print(f"Presentation saved as {output_ppt}")

if __name__ == '__main__':
    create_presentation()
